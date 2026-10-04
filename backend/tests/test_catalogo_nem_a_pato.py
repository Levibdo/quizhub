import csv
import json
import tempfile
import unittest
from io import BytesIO, StringIO
from pathlib import Path
from unittest.mock import patch
from uuid import uuid4

from openpyxl import Workbook
from sqlalchemy import create_engine, func, select
from sqlalchemy.orm import Session, sessionmaker
from sqlalchemy.pool import StaticPool

from app.cli import main as cli_main
from app.db.base import Base
from app.conteudo import CATEGORIAS_OFICIAIS, MODO_NEM_A_PATO
from app.models import Categoria, PerguntaNemPato
from app.services.catalogo_nem_a_pato import (
    ArquivoCatalogoNemAPatoInvalido,
    CatalogoNemAPatoInvalido,
    CatalogoPerguntasNemAPato,
    detectar_formato_catalogo,
)

CABECALHO = (
    "categoria_id",
    "enunciado",
    "resposta_numerica",
    "explicacao",
    "unidade",
    "fonte",
    "ativa",
)


def linha(**alteracoes):
    dados = {
        "categoria_id": "geral",
        "enunciado": "Quantos itens há?",
        "resposta_numerica": 42,
        "explicacao": "Contagem de teste.",
        "unidade": "itens",
        "fonte": "Fonte teste",
        "ativa": True,
    }
    dados.update(alteracoes)
    return dados


def criar_xlsx(registros, cabecalho=CABECALHO, *, formula=False, linhas_em_branco=False):
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.append(list(cabecalho))
    for registro in registros:
        if linhas_em_branco:
            worksheet.append([None] * len(cabecalho))
        if isinstance(registro, dict):
            worksheet.append([registro.get(campo) for campo in cabecalho])
        else:
            worksheet.append(list(registro))
        if formula:
            worksheet.cell(worksheet.max_row, cabecalho.index("resposta_numerica") + 1, "=40+2")
            formula = False
    arquivo = BytesIO()
    workbook.save(arquivo)
    workbook.close()
    return arquivo.getvalue()


def criar_csv(registros, cabecalho=CABECALHO, *, bom=False):
    arquivo = StringIO(newline="")
    escritor = csv.writer(arquivo, delimiter=",")
    escritor.writerow(cabecalho)
    for registro in registros:
        if isinstance(registro, dict):
            escritor.writerow([registro.get(campo) for campo in cabecalho])
        else:
            escritor.writerow(registro)
    conteudo = arquivo.getvalue().encode("utf-8")
    return b"\xef\xbb\xbf" + conteudo if bom else conteudo


def criar_json(registros):
    return json.dumps(registros, ensure_ascii=False).encode("utf-8")


class CatalogoNemAPatoTestCase(unittest.TestCase):
    def setUp(self):
        self.engine = create_engine(
            "sqlite://",
            connect_args={"check_same_thread": False},
            poolclass=StaticPool,
        )
        Base.metadata.create_all(self.engine)
        self.sessions = sessionmaker(bind=self.engine, expire_on_commit=False)
        with self.sessions() as session:
            session.add_all(
                [
                    Categoria(id=CATEGORIAS_OFICIAIS[MODO_NEM_A_PATO]["geral"], slug="geral", nome="Geral", modo=MODO_NEM_A_PATO, origem="OFICIAL", ativa=True),
                    Categoria(id=CATEGORIAS_OFICIAIS[MODO_NEM_A_PATO]["tecnologia"], slug="tecnologia", nome="Tecnologia", modo=MODO_NEM_A_PATO, origem="OFICIAL", ativa=True),
                    Categoria(id=uuid4(), slug="arquivada", nome="Arquivada", modo=MODO_NEM_A_PATO, origem="OFICIAL", ativa=False),
                ]
            )
            session.commit()
        self.service = CatalogoPerguntasNemAPato()

    def tearDown(self):
        self.engine.dispose()

    def count(self):
        with self.sessions() as session:
            return session.scalar(select(func.count()).select_from(PerguntaNemPato))

    def validar(self, conteudo, formato):
        with self.sessions() as session:
            return self.service.validar_arquivo(session, conteudo, formato)

    def importar(self, conteudo, formato):
        with self.sessions() as session:
            return self.service.importar_arquivo(session, conteudo, formato)

    def gravar(self, *, categoria="geral", enunciado="Já existe?", ativa=True):
        with self.sessions() as session:
            session.add(
                PerguntaNemPato(
                    categoria_id=(CATEGORIAS_OFICIAIS[MODO_NEM_A_PATO].get(categoria) or categoria),
                    enunciado=enunciado,
                    resposta_numerica=1,
                    explicacao="Explicação.",
                    ativa=ativa,
                )
            )
            session.commit()

    def cli(self, args, answers=()):
        iter_answers = iter(answers)
        output = []
        code = cli_main(
            args,
            input_fn=lambda _prompt: next(iter_answers),
            output=output.append,
            session_factory=self.sessions,
        )
        return code, output

    def arquivo_temporario(self, sufixo, conteudo):
        arquivo = tempfile.NamedTemporaryFile(suffix=sufixo, delete=False)
        self.addCleanup(Path(arquivo.name).unlink, missing_ok=True)
        arquivo.write(conteudo)
        arquivo.close()
        return arquivo.name


class TestCatalogoNemAPatoParsing(CatalogoNemAPatoTestCase):
    def test_xlsx_csv_json_validos(self):
        amostras = (
            (criar_xlsx([linha()]), "xlsx"),
            (criar_csv([linha()], bom=True), "csv"),
            (criar_json([linha()]), "json"),
        )
        for conteudo, formato in amostras:
            with self.subTest(formato=formato):
                relatorio = self.validar(conteudo, formato)
                self.assertEqual(
                    (relatorio.formato, relatorio.total, len(relatorio.validas), relatorio.invalidas),
                    (formato, 1, 1, 0),
                )

    def test_xlsx_ignora_linha_vazia_e_rejeita_formula_na_resposta(self):
        conteudo = criar_xlsx([linha(), linha(enunciado="Fórmula?")], formula=True, linhas_em_branco=True)
        relatorio = self.validar(conteudo, "xlsx")
        self.assertEqual((relatorio.total, len(relatorio.validas), relatorio.invalidas), (2, 1, 1))
        self.assertIn("não pode ser fórmula", relatorio.erros[0].motivo)

    def test_xlsx_preserva_inteiro_e_rejeita_fracao_e_booleano(self):
        relatorio = self.validar(
            criar_xlsx([
                linha(enunciado="Inteiro", resposta_numerica=40075),
                linha(enunciado="Fração", resposta_numerica=1.5),
                linha(enunciado="Booleano", resposta_numerica=True),
            ]),
            "xlsx",
        )
        self.assertEqual(relatorio.validas[0]["resposta_numerica"], 40075)
        self.assertEqual((len(relatorio.validas), relatorio.invalidas), (1, 2))

    def test_csv_utf8_com_acentos(self):
        relatorio = self.validar(
            criar_csv([linha(enunciado="Quantos quilômetros há?", explicacao="Explicação com acento.")]),
            "csv",
        )
        self.assertEqual(relatorio.validas[0]["enunciado"], "Quantos quilômetros há?")

    def test_xlsx_e_csv_sem_cabecalho_e_json_estruturalmente_invalido(self):
        with self.assertRaisesRegex(ArquivoCatalogoNemAPatoInvalido, "sem cabeçalho"):
            self.validar(b"", "csv")
        with self.assertRaisesRegex(ArquivoCatalogoNemAPatoInvalido, "sem cabeçalho"):
            self.validar(criar_xlsx([], cabecalho=()), "xlsx")
        with self.assertRaisesRegex(ArquivoCatalogoNemAPatoInvalido, "lista de objetos"):
            self.validar(json.dumps({"items": []}).encode(), "json")
        with self.assertRaisesRegex(ArquivoCatalogoNemAPatoInvalido, "malformado"):
            self.validar(b"[", "json")

    def test_json_lista_vazia_e_arquivo_vazio(self):
        relatorio = self.validar(b"[]", "json")
        self.assertEqual((relatorio.total, len(relatorio.validas)), (0, 0))
        with self.assertRaises(ArquivoCatalogoNemAPatoInvalido):
            self.validar(b"", "json")

    def test_cabecalho_ausente_e_extensao_nao_suportada(self):
        with self.assertRaisesRegex(ArquivoCatalogoNemAPatoInvalido, "obrigatórias"):
            self.validar(criar_csv([linha()], cabecalho=("enunciado",)), "csv")
        with self.assertRaisesRegex(ArquivoCatalogoNemAPatoInvalido, "extensão"):
            detectar_formato_catalogo("perguntas.xls")


class TestCatalogoNemAPatoValidacao(CatalogoNemAPatoTestCase):
    def test_categoria_enunciado_e_explicacao_validos(self):
        casos = (
            (linha(categoria_id="inexistente"), "categoria inexistente"),
            (linha(categoria_id="arquivada"), "categoria inativa"),
            (linha(enunciado=" \t "), "enunciado"),
            (linha(explicacao=" \t "), "explicacao"),
        )
        for registro, esperado in casos:
            with self.subTest(esperado=esperado):
                relatorio = self.validar(criar_json([registro]), "json")
                self.assertEqual(relatorio.invalidas, 1)
                self.assertIn(esperado, relatorio.erros[0].motivo)

    def test_valores_numericos_validos_e_invalidos(self):
        validos = (0, 42, 40075, "40075", 42.0)
        for valor in validos:
            with self.subTest(valor=valor):
                relatorio = self.validar(criar_json([linha(resposta_numerica=valor)]), "json")
                self.assertEqual(relatorio.invalidas, 0)
        invalidos = (-1, 1.5, "1,5", True, "abc", float("nan"), float("inf"), 2**63)
        for valor in invalidos:
            with self.subTest(valor=valor):
                relatorio = self.validar(criar_json([linha(resposta_numerica=valor)]), "json")
                self.assertEqual(relatorio.invalidas, 1)

    def test_booleano_ativa_opcionais_vazios_e_default(self):
        valores = {
            True: True,
            False: False,
            "true": True,
            "1": True,
            "sim": True,
            "S": True,
            "false": False,
            "0": False,
            "não": False,
            "N": False,
            None: True,
            " ": True,
        }
        for valor, esperado in valores.items():
            with self.subTest(valor=valor):
                relatorio = self.validar(
                    criar_json([linha(ativa=valor, unidade="  ", fonte=" ")]), "json"
                )
                self.assertEqual(relatorio.validas[0]["ativa"], esperado)
                self.assertIsNone(relatorio.validas[0]["unidade"])
                self.assertIsNone(relatorio.validas[0]["fonte"])
        for valor in (2, "yes", "sim talvez"):
            with self.subTest(invalida=valor):
                relatorio = self.validar(criar_json([linha(ativa=valor)]), "json")
                self.assertIn("ativa deve ser", relatorio.erros[0].motivo)

    def test_categoria_enunciado_explicacao_e_opcionais_sao_trimados(self):
        relatorio = self.validar(
            criar_json([linha(categoria_id=" geral ", enunciado="  Trim  ", explicacao=" Explica ", unidade=" km ", fonte=" F ")]),
            "json",
        )
        self.assertEqual(relatorio.validas[0], {
            "categoria_id": CATEGORIAS_OFICIAIS[MODO_NEM_A_PATO]["geral"],
            "enunciado": "Trim",
            "resposta_numerica": 42,
            "explicacao": "Explica",
            "unidade": "km",
            "fonte": "F",
            "ativa": True,
            "origem": "OFICIAL",
        })


class TestCatalogoNemAPatoDuplicatasImportacao(CatalogoNemAPatoTestCase):
    def test_duplicatas_no_arquivo_e_no_banco_sao_classificadas(self):
        self.gravar(enunciado="No banco?")
        relatorio = self.validar(
            criar_json([
                linha(enunciado="No banco?"),
                linha(enunciado="No arquivo?"),
                linha(enunciado="No arquivo?"),
            ]),
            "json",
        )
        self.assertEqual((len(relatorio.validas), relatorio.duplicadas, relatorio.invalidas), (1, 2, 0))

    def test_categoria_diferente_e_caixa_diferente_nao_sao_duplicatas(self):
        self.gravar(enunciado="Mesma pergunta?")
        relatorio = self.validar(
            criar_json([
                linha(categoria_id="tecnologia", enunciado="Mesma pergunta?"),
                linha(enunciado="mesma pergunta?"),
            ]),
            "json",
        )
        self.assertEqual((len(relatorio.validas), relatorio.duplicadas), (2, 0))

    def test_importacao_valida_ignora_duplicatas_e_reimportacao_e_idempotente(self):
        conteudo = criar_csv([linha(enunciado="Importada?"), linha(enunciado="Importada?")])
        resultado = self.importar(conteudo, "csv")
        segunda = self.importar(conteudo, "csv")
        self.assertEqual((resultado.criadas, resultado.duplicadas, resultado.falhas), (1, 1, 0))
        self.assertEqual((segunda.criadas, segunda.duplicadas, segunda.falhas), (0, 2, 0))
        self.assertEqual(self.count(), 1)

    def test_qualquer_invalida_cancela_importacao_sem_parcial(self):
        conteudo = criar_json([linha(enunciado="Válida?"), linha(resposta_numerica=-1)])
        with self.sessions() as session:
            with self.assertRaises(CatalogoNemAPatoInvalido):
                self.service.importar_arquivo(session, conteudo, "json")
        self.assertEqual(self.count(), 0)

    def test_erro_inesperado_durante_flush_faz_rollback(self):
        conteudo = criar_json([linha(enunciado="Rollback?")])
        with self.sessions() as session:
            with patch.object(session, "flush", side_effect=RuntimeError("falha simulada")):
                with self.assertRaisesRegex(RuntimeError, "falha simulada"):
                    self.service.importar_arquivo(session, conteudo, "json")
        self.assertEqual(self.count(), 0)


class TestCatalogoNemAPatoCli(CatalogoNemAPatoTestCase):
    def executar_arquivo(self, comando, registros, *, respostas=(), extensao=".json"):
        caminho = self.arquivo_temporario(extensao, criar_json(registros))
        return self.cli(["nem-pato", "perguntas", comando, caminho], respostas)

    def test_validar_nao_altera_banco(self):
        codigo, saidas = self.executar_arquivo("validar", [linha()])
        self.assertEqual(codigo, 0)
        self.assertEqual(self.count(), 0)
        self.assertIn("Formato: json | Total: 1 | Válidas: 1 | Duplicadas: 0 | Inválidas: 0", saidas[0])

    def test_confirmacao_negativa_nao_importa_e_positiva_importa(self):
        caminho = self.arquivo_temporario(".json", criar_json([linha()]))
        codigo, saidas = self.cli(["nem-pato", "perguntas", "importar", caminho], ("n",))
        self.assertEqual(codigo, 0)
        self.assertEqual(self.count(), 0)
        self.assertTrue(any("cancelada" in saida for saida in saidas))
        codigo, saidas = self.cli(["nem-pato", "perguntas", "importar", caminho], ("s",))
        self.assertEqual(codigo, 0)
        self.assertEqual(self.count(), 1)
        self.assertIn("1 criada(s), 0 duplicada(s), 0 falha(s)", saidas[-1])

    def test_importar_invalida_nao_pede_confirmacao_nem_grava(self):
        caminho = self.arquivo_temporario(".json", criar_json([linha(resposta_numerica=-1)]))
        codigo, saidas = self.cli(["nem-pato", "perguntas", "importar", caminho])
        self.assertEqual(codigo, 1)
        self.assertEqual(self.count(), 0)
        self.assertTrue(any("há registros inválidos" in saida for saida in saidas))

    def test_listar_resumo_e_criar(self):
        caminho = self.arquivo_temporario(
            ".json", criar_json([linha(), linha(enunciado="Inativa?", ativa=False)])
        )
        self.cli(["nem-pato", "perguntas", "importar", caminho], ("sim",))
        _, listagem = self.cli(["nem-pato", "perguntas", "listar"])
        self.assertTrue(any("Quantos itens há?" in saida for saida in listagem))
        self.assertFalse(any("42" in saida for saida in listagem))
        _, resumo = self.cli(["nem-pato", "perguntas", "resumo"])
        self.assertIn("Total: 2 | Ativas: 1 | Inativas: 1", resumo[0])
        self.assertIn("geral: 2", resumo)
        entradas = ("geral", "Nova pergunta?", "40075", "km", "Explicação nova.", "Fonte", "s")
        codigo, saidas = self.cli(["nem-pato", "perguntas", "criar"], entradas)
        self.assertEqual(codigo, 0)
        self.assertTrue(any("criada com sucesso" in saida for saida in saidas))
        duplicada, erros = self.cli(["nem-pato", "perguntas", "criar"], entradas)
        self.assertEqual(duplicada, 1)
        self.assertTrue(any("duplicada" in saida for saida in erros))


if __name__ == "__main__":
    unittest.main()
