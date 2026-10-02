import csv
import json
import math
import re
from dataclasses import dataclass
from io import BytesIO, StringIO
from pathlib import Path
from zipfile import BadZipFile

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from sqlalchemy import func, select, text
from sqlalchemy.orm import Session

from app.models import Categoria, PerguntaNemPato

COLUNAS_OBRIGATORIAS = (
    "categoria_id",
    "enunciado",
    "resposta_numerica",
    "explicacao",
)
COLUNAS_OPCIONAIS = ("unidade", "fonte", "ativa")
COLUNAS_CANONICAS = COLUNAS_OBRIGATORIAS + COLUNAS_OPCIONAIS
FORMATOS_SUPORTADOS = {"xlsx", "csv", "json"}
BIGINT_MAX = 2**63 - 1
_INTEGER_TEXT = re.compile(r"^[0-9]+$")


class ArquivoCatalogoNemAPatoInvalido(ValueError):
    """Arquivo malformado ou sem a estrutura canônica esperada."""


class CatalogoNemAPatoInvalido(ValueError):
    """Uma ou mais linhas não passaram na validação do catálogo."""

    def __init__(self, relatorio: "RelatorioCatalogoNemAPato"):
        self.relatorio = relatorio
        super().__init__("há registros inválidos; nenhum registro foi importado")


@dataclass(frozen=True)
class CandidataNemAPato:
    referencia: int
    dados: dict[str, object]
    erro_estrutura: str | None = None


@dataclass(frozen=True)
class ErroCatalogoNemAPato:
    referencia: int
    motivo: str


@dataclass(frozen=True)
class RelatorioCatalogoNemAPato:
    formato: str
    total: int
    validas: tuple[dict[str, object], ...]
    duplicadas: int
    invalidas: int
    erros: tuple[ErroCatalogoNemAPato, ...]


@dataclass(frozen=True)
class ResultadoImportacaoNemAPato:
    criadas: int
    duplicadas: int
    falhas: int


def detectar_formato_catalogo(caminho: str | Path) -> str:
    formato = Path(caminho).suffix.lower().removeprefix(".")
    if formato not in FORMATOS_SUPORTADOS:
        raise ArquivoCatalogoNemAPatoInvalido(
            "extensão não suportada; use .xlsx, .csv ou .json"
        )
    return formato


def _linha_vazia(valores) -> bool:
    return all(
        valor is None or (isinstance(valor, str) and not valor.strip())
        for valor in valores
    )


def _indices_cabecalho(cabecalho) -> dict[str, int]:
    if cabecalho is None:
        raise ArquivoCatalogoNemAPatoInvalido("arquivo sem cabeçalho")
    nomes = [valor.strip() if isinstance(valor, str) else "" for valor in cabecalho]
    repetidas = sorted({nome for nome in nomes if nome and nomes.count(nome) > 1})
    if repetidas:
        raise ArquivoCatalogoNemAPatoInvalido(
            "colunas duplicadas: " + ", ".join(repetidas)
        )
    ausentes = [nome for nome in COLUNAS_OBRIGATORIAS if nome not in nomes]
    if ausentes:
        raise ArquivoCatalogoNemAPatoInvalido(
            "colunas obrigatórias ausentes: " + ", ".join(ausentes)
        )
    return {nome: nomes.index(nome) for nome in COLUNAS_CANONICAS if nome in nomes}


def _dados_por_cabecalho(valores, indices) -> dict[str, object]:
    return {
        nome: valores[indice] if indice < len(valores) else None
        for nome, indice in indices.items()
    }


def ler_catalogo_xlsx(conteudo: bytes) -> list[CandidataNemAPato]:
    try:
        workbook = load_workbook(BytesIO(conteudo), read_only=True, data_only=False)
    except (BadZipFile, InvalidFileException, KeyError, OSError) as erro:
        raise ArquivoCatalogoNemAPatoInvalido("arquivo XLSX inválido") from erro
    try:
        linhas = workbook.active.iter_rows()
        primeira = next(linhas, None)
        indices = _indices_cabecalho(
            [celula.value for celula in primeira] if primeira is not None else None
        )
        candidatas = []
        for numero, celulas in enumerate(linhas, start=2):
            valores = [celula.value for celula in celulas]
            if _linha_vazia(valores):
                continue
            dados = _dados_por_cabecalho(valores, indices)
            indice_resposta = indices["resposta_numerica"]
            erro = (
                "resposta_numerica não pode ser fórmula"
                if indice_resposta < len(celulas)
                and celulas[indice_resposta].data_type == "f"
                else None
            )
            candidatas.append(CandidataNemAPato(numero, dados, erro))
        return candidatas
    finally:
        workbook.close()


def ler_catalogo_csv(conteudo: bytes) -> list[CandidataNemAPato]:
    try:
        texto_csv = conteudo.decode("utf-8-sig")
    except UnicodeDecodeError as erro:
        raise ArquivoCatalogoNemAPatoInvalido("CSV deve usar UTF-8") from erro
    try:
        leitor = csv.reader(StringIO(texto_csv, newline=""), delimiter=",", strict=True)
        indices = _indices_cabecalho(next(leitor, None))
        candidatas = []
        for valores in leitor:
            referencia = leitor.line_num
            if _linha_vazia(valores):
                continue
            candidatas.append(
                CandidataNemAPato(
                    referencia, _dados_por_cabecalho(valores, indices)
                )
            )
        return candidatas
    except csv.Error as erro:
        raise ArquivoCatalogoNemAPatoInvalido(f"CSV inválido: {erro}") from erro


def ler_catalogo_json(conteudo: bytes) -> list[CandidataNemAPato]:
    try:
        registros = json.loads(conteudo.decode("utf-8-sig"))
    except (UnicodeDecodeError, json.JSONDecodeError) as erro:
        raise ArquivoCatalogoNemAPatoInvalido("JSON malformado ou fora de UTF-8") from erro
    if not isinstance(registros, list):
        raise ArquivoCatalogoNemAPatoInvalido(
            "JSON deve conter uma lista de objetos no nível superior"
        )
    candidatas = []
    for indice, registro in enumerate(registros, start=1):
        if not isinstance(registro, dict):
            candidatas.append(
                CandidataNemAPato(indice, {}, "item JSON deve ser um objeto")
            )
            continue
        candidatas.append(
            CandidataNemAPato(
                indice,
                {nome: registro.get(nome) for nome in COLUNAS_CANONICAS},
            )
        )
    return candidatas


LEITORES_CATALOGO = {
    "xlsx": ler_catalogo_xlsx,
    "csv": ler_catalogo_csv,
    "json": ler_catalogo_json,
}


def _texto_obrigatorio(valor: object, campo: str) -> tuple[str | None, str | None]:
    if not isinstance(valor, str):
        return None, f"campo obrigatório deve ser texto: {campo}"
    normalizado = valor.strip()
    if not normalizado:
        return None, f"campo obrigatório vazio: {campo}"
    return normalizado, None


def _resposta_inteira(valor: object) -> tuple[int | None, str | None]:
    if isinstance(valor, bool):
        return None, "resposta_numerica deve ser inteiro não negativo"
    if isinstance(valor, int):
        numero = valor
    elif isinstance(valor, float):
        if not math.isfinite(valor) or not valor.is_integer():
            return None, "resposta_numerica deve ser inteiro não negativo"
        numero = int(valor)
        # Excel só representa exatamente inteiros até 2**53 - 1 como float.
        if numero > 2**53 - 1:
            return None, "resposta_numerica excede a precisão inteira segura do XLSX"
    elif isinstance(valor, str) and _INTEGER_TEXT.fullmatch(valor.strip()):
        numero = int(valor.strip())
    else:
        return None, "resposta_numerica deve ser inteiro não negativo"
    if numero < 0:
        return None, "resposta_numerica deve ser inteiro não negativo"
    if numero > BIGINT_MAX:
        return None, "resposta_numerica excede o limite de BIGINT"
    return numero, None


def _ativa(valor: object) -> tuple[bool | None, str | None]:
    if valor is None or (isinstance(valor, str) and not valor.strip()):
        return True, None
    if isinstance(valor, bool):
        return valor, None
    if isinstance(valor, int) and valor in (0, 1):
        return bool(valor), None
    if isinstance(valor, str):
        normalizado = valor.strip().casefold()
        valores = {
            "true": True,
            "1": True,
            "sim": True,
            "s": True,
            "false": False,
            "0": False,
            "nao": False,
            "não": False,
            "n": False,
        }
        if normalizado in valores:
            return valores[normalizado], None
    return None, "ativa deve ser booleano, true/false, 1/0, sim/não ou s/n"


def validar_registro_nem_a_pato(
    dados: dict[str, object], categorias: dict[str, Categoria]
) -> tuple[dict[str, object] | None, str | None]:
    categoria_id, erro = _texto_obrigatorio(dados.get("categoria_id"), "categoria_id")
    if erro:
        return None, erro
    enunciado, erro = _texto_obrigatorio(dados.get("enunciado"), "enunciado")
    if erro:
        return None, erro
    explicacao, erro = _texto_obrigatorio(dados.get("explicacao"), "explicacao")
    if erro:
        return None, erro
    categoria = categorias.get(categoria_id)
    if categoria is None:
        return None, "categoria inexistente"
    if not categoria.ativa:
        return None, "categoria inativa"
    resposta, erro = _resposta_inteira(dados.get("resposta_numerica"))
    if erro:
        return None, erro
    ativa, erro = _ativa(dados.get("ativa"))
    if erro:
        return None, erro
    opcionais = {}
    for campo in ("unidade", "fonte"):
        valor = dados.get(campo)
        if valor is not None and not isinstance(valor, str):
            return None, f"campo opcional deve ser texto: {campo}"
        opcionais[campo] = valor.strip() or None if isinstance(valor, str) else None
    return {
        "categoria_id": categoria_id,
        "enunciado": enunciado,
        "resposta_numerica": resposta,
        "explicacao": explicacao,
        **opcionais,
        "ativa": ativa,
    }, None


class CatalogoPerguntasNemAPato:
    def ler(self, conteudo: bytes, formato: str) -> list[CandidataNemAPato]:
        formato = formato.lower().removeprefix(".")
        leitor = LEITORES_CATALOGO.get(formato)
        if leitor is None:
            raise ArquivoCatalogoNemAPatoInvalido(
                "formato não suportado; use xlsx, csv ou json"
            )
        return leitor(conteudo)

    def validar_arquivo(
        self, db: Session, conteudo: bytes, formato: str
    ) -> RelatorioCatalogoNemAPato:
        return self._auditar(db, self.ler(conteudo, formato), formato)

    def validar_registro(
        self, db: Session, dados: dict[str, object]
    ) -> tuple[dict[str, object] | None, str | None, bool]:
        categorias = {
            categoria.id: categoria
            for categoria in db.scalars(select(Categoria))
        }
        validado, erro = validar_registro_nem_a_pato(dados, categorias)
        if erro:
            return None, erro, False
        duplicada = self._chave(validado) in self._chaves_existentes(db)
        if duplicada:
            return None, "pergunta duplicada: categoria_id e enunciado já existentes", True
        return validado, None, False

    def criar(self, db: Session, dados: dict[str, object]) -> PerguntaNemPato:
        try:
            if db.bind is not None and db.bind.dialect.name == "postgresql":
                db.rollback()
                db.execute(
                    text(
                        "LOCK TABLE perguntas_nem_pato "
                        "IN SHARE ROW EXCLUSIVE MODE"
                    )
                )
            validado, erro, duplicada = self.validar_registro(db, dados)
            if erro:
                prefixo = "duplicada" if duplicada else "inválida"
                raise ValueError(f"pergunta {prefixo}: {erro}")
            pergunta = PerguntaNemPato(**validado)
            db.add(pergunta)
            db.flush()
            db.commit()
            return pergunta
        except Exception:
            db.rollback()
            raise

    def importar_arquivo(
        self, db: Session, conteudo: bytes, formato: str
    ) -> ResultadoImportacaoNemAPato:
        candidatas = self.ler(conteudo, formato)
        try:
            # Serializa importações concorrentes do catálogo no PostgreSQL.
            if db.bind is not None and db.bind.dialect.name == "postgresql":
                db.rollback()
                db.execute(
                    text(
                        "LOCK TABLE perguntas_nem_pato "
                        "IN SHARE ROW EXCLUSIVE MODE"
                    )
                )
            auditoria = self._auditar(db, candidatas, formato)
            if auditoria.invalidas:
                raise CatalogoNemAPatoInvalido(auditoria)
            for dados in auditoria.validas:
                db.add(PerguntaNemPato(**dados))
            db.flush()
            db.commit()
            return ResultadoImportacaoNemAPato(
                criadas=len(auditoria.validas),
                duplicadas=auditoria.duplicadas,
                falhas=0,
            )
        except Exception:
            db.rollback()
            raise

    def listar(self, db: Session) -> list[tuple[int, str, str, str | None, bool]]:
        linhas = db.execute(
            select(
                PerguntaNemPato.id,
                PerguntaNemPato.categoria_id,
                PerguntaNemPato.enunciado,
                PerguntaNemPato.unidade,
                PerguntaNemPato.ativa,
            ).order_by(PerguntaNemPato.id)
        )
        return list(linhas.tuples())

    def resumo(self, db: Session) -> dict[str, object]:
        total, ativas, inativas = db.execute(
            select(
                func.count(PerguntaNemPato.id),
                func.count(PerguntaNemPato.id).filter(PerguntaNemPato.ativa.is_(True)),
                func.count(PerguntaNemPato.id).filter(PerguntaNemPato.ativa.is_(False)),
            )
        ).one()
        por_categoria = db.execute(
            select(
                PerguntaNemPato.categoria_id,
                func.count(PerguntaNemPato.id),
            )
            .group_by(PerguntaNemPato.categoria_id)
            .order_by(PerguntaNemPato.categoria_id)
        ).all()
        return {
            "total": total,
            "ativas": ativas,
            "inativas": inativas,
            "por_categoria": list(por_categoria),
        }

    @staticmethod
    def _chave(dados: dict[str, object]) -> tuple[str, str]:
        return str(dados["categoria_id"]), str(dados["enunciado"]).strip()

    @staticmethod
    def _chaves_existentes(db: Session) -> set[tuple[str, str]]:
        existentes = db.execute(
            select(PerguntaNemPato.categoria_id, PerguntaNemPato.enunciado)
        ).tuples()
        return {
            (categoria_id, enunciado.strip())
            for categoria_id, enunciado in existentes
        }

    def _auditar(
        self, db: Session, candidatas: list[CandidataNemAPato], formato: str
    ) -> RelatorioCatalogoNemAPato:
        categorias = {
            categoria.id: categoria
            for categoria in db.scalars(select(Categoria))
        }
        chaves_conhecidas = self._chaves_existentes(db)
        validas = []
        erros = []
        duplicadas = 0
        invalidas = 0
        for candidata in candidatas:
            if candidata.erro_estrutura:
                erros.append(
                    ErroCatalogoNemAPato(candidata.referencia, candidata.erro_estrutura)
                )
                invalidas += 1
                continue
            dados, motivo = validar_registro_nem_a_pato(candidata.dados, categorias)
            if motivo:
                erros.append(ErroCatalogoNemAPato(candidata.referencia, motivo))
                invalidas += 1
                continue
            chave = self._chave(dados)
            if chave in chaves_conhecidas:
                duplicadas += 1
                erros.append(
                    ErroCatalogoNemAPato(
                        candidata.referencia,
                        "pergunta duplicada: categoria_id e enunciado já existentes",
                    )
                )
                continue
            chaves_conhecidas.add(chave)
            validas.append(dados)
        return RelatorioCatalogoNemAPato(
            formato=formato.lower().removeprefix("."),
            total=len(candidatas),
            validas=tuple(validas),
            duplicadas=duplicadas,
            invalidas=invalidas,
            erros=tuple(erros),
        )


catalogo_perguntas_nem_a_pato = CatalogoPerguntasNemAPato()