import csv
import hashlib
import json
import math
import re
from dataclasses import dataclass
from datetime import datetime, timedelta, timezone
from io import BytesIO, StringIO
from pathlib import Path
from typing import Any
from uuid import UUID
from zipfile import BadZipFile

import jwt
from jwt.exceptions import ExpiredSignatureError, InvalidTokenError
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from sqlalchemy import func, select, text
from sqlalchemy.orm import Session

from app.conteudo import (
    LIMITE_PERGUNTAS_USUARIO_POR_MODO,
    MODO_NEM_A_PATO,
    MODO_QUIZ_CLASSICO,
    ORIGEM_USUARIO,
)
from app.models import Categoria, Pergunta, PerguntaNemPato, Usuario
from app.schemas.meu_conteudo_importacoes import (
    ArquivoImportacaoPreview,
    ErroImportacaoPrivada,
    PreviewImportacaoPrivada,
    QuotaImportacaoPreview,
    ReferenciaErroImportacao,
    ResultadoImportacaoPrivada,
)
from app.security import JWT_ALGORITHM, obter_segredo_jwt


MAX_ARQUIVO = 5 * 1024 * 1024
MAX_REGISTROS = 100
MAX_TEXTO_CAMPO = 10_000
MAX_TEXTO_TOTAL = 1024 * 1024
BIGINT_MAX = 2**63 - 1
EXPIRACAO_PREVIEW = timedelta(minutes=15)
VERSAO_CONTRATO = 1
FINALIDADE_TOKEN = "importacao_privada_preview"
FORMATOS = {"xlsx", "csv", "json"}
INTEIRO_CSV = re.compile(r"^[0-9]+$")

COLUNAS_CLASSIC = (
    "categoria_id", "enunciado", "alternativa_a", "alternativa_b",
    "alternativa_c", "alternativa_d", "alternativa_correta", "explicacao",
)
COLUNAS_NP_OBRIGATORIAS = (
    "categoria_id", "enunciado", "resposta_numerica", "explicacao",
)
COLUNAS_NP_OPCIONAIS = ("unidade", "fonte")
COLUNAS_NP = COLUNAS_NP_OBRIGATORIAS + COLUNAS_NP_OPCIONAIS
ALTERNATIVAS = {"A": 0, "B": 1, "C": 2, "D": 3}


class ArquivoImportacaoPrivadaInvalido(ValueError):
    pass


class TokenPreviewInvalido(ValueError):
    pass


class ConflitoImportacaoPrivada(Exception):
    detalhe = "o lote não é mais confirmável; valide o arquivo novamente"


class QuotaImportacaoPrivadaAlterada(ConflitoImportacaoPrivada):
    detalhe = "a quota disponível mudou; valide o arquivo novamente"


@dataclass(frozen=True)
class Candidata:
    referencia: int
    tipo_referencia: str
    dados: dict[str, Any]
    erro_campo: str | None = None
    erro_codigo: str | None = None
    erro_mensagem: str | None = None


@dataclass(frozen=True)
class Auditoria:
    validas: list[dict[str, Any]]
    erros: list[ErroImportacaoPrivada]
    total: int


def formato_do_nome(nome: str | None) -> str:
    formato = Path(nome or "").suffix.lower().removeprefix(".")
    if formato not in FORMATOS:
        raise ArquivoImportacaoPrivadaInvalido(
            "extensão não suportada; use .xlsx, .csv ou .json"
        )
    return formato


def _erro(candidata: Candidata, campo: str | None, codigo: str, mensagem: str):
    return ErroImportacaoPrivada(
        referencia=ReferenciaErroImportacao(
            tipo=candidata.tipo_referencia, valor=candidata.referencia
        ),
        campo=campo,
        codigo=codigo,
        mensagem=mensagem,
    )


def _cabecalho(valores, modo: str) -> tuple[list[str], dict[str, int]]:
    if valores is None:
        raise ArquivoImportacaoPrivadaInvalido("arquivo sem cabeçalho")
    nomes = [valor.strip() if isinstance(valor, str) else "" for valor in valores]
    if not any(nomes):
        raise ArquivoImportacaoPrivadaInvalido("arquivo sem cabeçalho")
    repetidas = sorted({nome for nome in nomes if nome and nomes.count(nome) > 1})
    if repetidas:
        raise ArquivoImportacaoPrivadaInvalido(
            "colunas duplicadas: " + ", ".join(repetidas)
        )
    permitidas = COLUNAS_CLASSIC if modo == MODO_QUIZ_CLASSICO else COLUNAS_NP
    obrigatorias = (
        COLUNAS_CLASSIC
        if modo == MODO_QUIZ_CLASSICO
        else COLUNAS_NP_OBRIGATORIAS
    )
    inesperadas = sorted(nome for nome in nomes if nome not in permitidas)
    if inesperadas:
        raise ArquivoImportacaoPrivadaInvalido(
            "colunas inesperadas: " + ", ".join(inesperadas)
        )
    ausentes = [nome for nome in obrigatorias if nome not in nomes]
    if ausentes:
        raise ArquivoImportacaoPrivadaInvalido(
            "colunas obrigatórias ausentes: " + ", ".join(ausentes)
        )
    return nomes, {nome: nomes.index(nome) for nome in permitidas if nome in nomes}


def _linha_vazia(valores) -> bool:
    return all(
        valor is None or (isinstance(valor, str) and not valor.strip())
        for valor in valores
    )


def _limitar_quantidade(candidatas: list[Candidata]) -> None:
    if not candidatas:
        raise ArquivoImportacaoPrivadaInvalido("arquivo não contém registros")
    if len(candidatas) > MAX_REGISTROS:
        raise ArquivoImportacaoPrivadaInvalido(
            f"o lote deve conter no máximo {MAX_REGISTROS} registros"
        )


def _ler_xlsx(conteudo: bytes, modo: str) -> list[Candidata]:
    try:
        workbook = load_workbook(
            BytesIO(conteudo), read_only=True, data_only=False,
            keep_links=False,
        )
    except (BadZipFile, InvalidFileException, KeyError, OSError, ValueError) as erro:
        raise ArquivoImportacaoPrivadaInvalido("arquivo XLSX inválido") from erro
    try:
        linhas = workbook.active.iter_rows()
        primeira = next(linhas, None)
        _, indices = _cabecalho(
            [celula.value for celula in primeira] if primeira else None, modo
        )
        candidatas: list[Candidata] = []
        for numero, celulas in enumerate(linhas, start=2):
            valores = [celula.value for celula in celulas]
            if _linha_vazia(valores):
                continue
            dados = {
                nome: valores[indice] if indice < len(valores) else None
                for nome, indice in indices.items()
            }
            formula = next(
                (
                    nome for nome, indice in indices.items()
                    if indice < len(celulas) and celulas[indice].data_type == "f"
                ),
                None,
            )
            candidatas.append(Candidata(
                numero, "linha", dados,
                formula,
                "formula_nao_permitida" if formula else None,
                "fórmulas não são permitidas" if formula else None,
            ))
            if len(candidatas) > MAX_REGISTROS:
                break
        _limitar_quantidade(candidatas)
        return candidatas
    finally:
        workbook.close()


def _ler_csv(conteudo: bytes, modo: str) -> list[Candidata]:
    try:
        texto_csv = conteudo.decode("utf-8-sig")
    except UnicodeDecodeError as erro:
        raise ArquivoImportacaoPrivadaInvalido("CSV deve usar UTF-8") from erro
    try:
        leitor = csv.reader(StringIO(texto_csv, newline=""), delimiter=",", strict=True)
        nomes, indices = _cabecalho(next(leitor, None), modo)
        candidatas = []
        for valores in leitor:
            if _linha_vazia(valores):
                continue
            if len(valores) > len(nomes):
                candidatas.append(Candidata(
                    leitor.line_num, "linha", {}, None, "estrutura_invalida",
                    "a linha possui colunas além do cabeçalho",
                ))
                continue
            candidatas.append(Candidata(
                leitor.line_num,
                "linha",
                {
                    nome: valores[indice] if indice < len(valores) else None
                    for nome, indice in indices.items()
                },
            ))
            if len(candidatas) > MAX_REGISTROS:
                break
        _limitar_quantidade(candidatas)
        return candidatas
    except csv.Error as erro:
        raise ArquivoImportacaoPrivadaInvalido(f"CSV inválido: {erro}") from erro


def _objeto_json_sem_chaves_repetidas(pares):
    resultado = {}
    for chave, valor in pares:
        if chave in resultado:
            raise ArquivoImportacaoPrivadaInvalido(
                f"JSON contém chave duplicada: {chave}"
            )
        resultado[chave] = valor
    return resultado


def _ler_json(conteudo: bytes, modo: str) -> list[Candidata]:
    try:
        documento = json.loads(
            conteudo.decode("utf-8-sig"),
            object_pairs_hook=_objeto_json_sem_chaves_repetidas,
        )
    except ArquivoImportacaoPrivadaInvalido:
        raise
    except (UnicodeDecodeError, json.JSONDecodeError, RecursionError) as erro:
        raise ArquivoImportacaoPrivadaInvalido(
            "JSON malformado ou fora de UTF-8"
        ) from erro
    if not isinstance(documento, list):
        raise ArquivoImportacaoPrivadaInvalido(
            "JSON deve conter uma lista no nível superior"
        )
    if not documento:
        raise ArquivoImportacaoPrivadaInvalido("arquivo não contém registros")
    if len(documento) > MAX_REGISTROS:
        raise ArquivoImportacaoPrivadaInvalido(
            f"o lote deve conter no máximo {MAX_REGISTROS} registros"
        )
    permitidas = set(COLUNAS_CLASSIC if modo == MODO_QUIZ_CLASSICO else COLUNAS_NP)
    candidatas = []
    for indice, item in enumerate(documento, start=1):
        if not isinstance(item, dict):
            candidatas.append(Candidata(
                indice, "item", {}, None, "estrutura_invalida",
                "o item deve ser um objeto",
            ))
            continue
        inesperadas = sorted(set(item) - permitidas)
        candidatas.append(Candidata(
            indice, "item", item,
            inesperadas[0] if inesperadas else None,
            "campo_inesperado" if inesperadas else None,
            f"campo inesperado: {inesperadas[0]}" if inesperadas else None,
        ))
    return candidatas


def ler(conteudo: bytes, formato: str, modo: str) -> list[Candidata]:
    if not conteudo:
        raise ArquivoImportacaoPrivadaInvalido("arquivo vazio")
    leitores = {"xlsx": _ler_xlsx, "csv": _ler_csv, "json": _ler_json}
    return leitores[formato](conteudo, modo)


def _texto(candidata, campo, *, opcional=False):
    valor = candidata.dados.get(campo)
    if opcional and (valor is None or (isinstance(valor, str) and not valor.strip())):
        return None, None
    if not isinstance(valor, str):
        return None, _erro(
            candidata, campo, "tipo_invalido", "deve ser texto"
        )
    if len(valor) > MAX_TEXTO_CAMPO:
        return None, _erro(
            candidata, campo, "texto_muito_longo",
            f"deve ter no máximo {MAX_TEXTO_CAMPO} caracteres",
        )
    normalizado = valor.strip()
    if not normalizado and not opcional:
        return None, _erro(candidata, campo, "campo_obrigatorio", "não pode ser vazio")
    return normalizado or None, None


def _categoria_id(candidata):
    valor = candidata.dados.get("categoria_id")
    if not isinstance(valor, str):
        return None, _erro(candidata, "categoria_id", "uuid_invalido", "deve ser um UUID")
    try:
        return UUID(valor.strip()), None
    except (ValueError, AttributeError):
        return None, _erro(candidata, "categoria_id", "uuid_invalido", "deve ser um UUID")


def _numero_np(candidata: Candidata, formato: str):
    valor = candidata.dados.get("resposta_numerica")
    numero = None
    if isinstance(valor, bool):
        pass
    elif formato == "json":
        if isinstance(valor, int):
            numero = valor
    elif formato == "csv":
        if isinstance(valor, str) and INTEIRO_CSV.fullmatch(valor):
            numero = int(valor)
    elif isinstance(valor, int):
        numero = valor
    elif isinstance(valor, float) and math.isfinite(valor) and valor.is_integer():
        if valor <= 2**53 - 1:
            numero = int(valor)
        else:
            return None, _erro(
                candidata, "resposta_numerica", "precisao_insegura",
                "excede a precisão inteira segura do XLSX",
            )
    if numero is None:
        return None, _erro(
            candidata, "resposta_numerica", "inteiro_invalido",
            "deve ser um inteiro não negativo",
        )
    if numero < 0 or numero > BIGINT_MAX:
        return None, _erro(
            candidata, "resposta_numerica", "fora_da_faixa",
            f"deve estar entre 0 e {BIGINT_MAX}",
        )
    return numero, None


def _validar_campos(candidata: Candidata, formato: str, modo: str):
    if candidata.erro_codigo:
        return None, _erro(
            candidata, candidata.erro_campo, candidata.erro_codigo,
            candidata.erro_mensagem or "registro inválido",
        )
    categoria_id, erro = _categoria_id(candidata)
    if erro:
        return None, erro
    enunciado, erro = _texto(candidata, "enunciado")
    if erro:
        return None, erro
    if modo == MODO_QUIZ_CLASSICO:
        dados = {"categoria_id": categoria_id, "enunciado": enunciado}
        for campo in (
            "alternativa_a", "alternativa_b", "alternativa_c", "alternativa_d",
            "explicacao",
        ):
            dados[campo], erro = _texto(candidata, campo)
            if erro:
                return None, erro
        alternativa = candidata.dados.get("alternativa_correta")
        if not isinstance(alternativa, str) or alternativa.strip().upper() not in ALTERNATIVAS:
            return None, _erro(
                candidata, "alternativa_correta", "alternativa_invalida",
                "deve ser A, B, C ou D",
            )
        dados["alternativa_correta"] = ALTERNATIVAS[alternativa.strip().upper()]
        return dados, None
    resposta, erro = _numero_np(candidata, formato)
    if erro:
        return None, erro
    explicacao, erro = _texto(candidata, "explicacao")
    if erro:
        return None, erro
    unidade, erro = _texto(candidata, "unidade", opcional=True)
    if erro:
        return None, erro
    fonte, erro = _texto(candidata, "fonte", opcional=True)
    if erro:
        return None, erro
    return {
        "categoria_id": categoria_id,
        "enunciado": enunciado,
        "resposta_numerica": resposta,
        "explicacao": explicacao,
        "unidade": unidade,
        "fonte": fonte,
    }, None


def _categorias_validas(
    db: Session, usuario_id: UUID, modo: str, ids: set[UUID], *, bloquear=False
) -> dict[UUID, Categoria]:
    consulta = select(Categoria).where(
        Categoria.id.in_(sorted(ids, key=str)),
        Categoria.origem == ORIGEM_USUARIO,
        Categoria.usuario_id == usuario_id,
        Categoria.modo == modo,
        Categoria.ativa.is_(True),
        Categoria.excluida_em.is_(None),
    ).order_by(Categoria.id)
    if bloquear:
        consulta = consulta.with_for_update()
    return {categoria.id: categoria for categoria in db.scalars(consulta)}


def _modelo(modo: str):
    return Pergunta if modo == MODO_QUIZ_CLASSICO else PerguntaNemPato


def _auditar(
    db: Session, usuario_id: UUID, modo: str, formato: str,
    candidatas: list[Candidata], *, bloquear_categorias=False,
) -> Auditoria:
    preliminares = []
    erros = []
    total_texto = 0
    for candidata in candidatas:
        dados, erro = _validar_campos(candidata, formato, modo)
        if erro:
            erros.append(erro)
            continue
        total_texto += sum(
            len(valor) for valor in dados.values() if isinstance(valor, str)
        )
        if total_texto > MAX_TEXTO_TOTAL:
            erros.append(_erro(
                candidata, None, "texto_total_excedido",
                "o lote excede 1 MiB de texto normalizado",
            ))
            continue
        preliminares.append((candidata, dados))

    ids = {dados["categoria_id"] for _, dados in preliminares}
    categorias = _categorias_validas(
        db, usuario_id, modo, ids, bloquear=bloquear_categorias
    ) if ids else {}
    modelo = _modelo(modo)
    existentes = set(db.execute(select(
        modelo.categoria_id, modelo.enunciado
    ).where(
        modelo.origem == ORIGEM_USUARIO,
        modelo.usuario_id == usuario_id,
        modelo.excluida_em.is_(None),
        modelo.categoria_id.in_(ids),
    )).tuples()) if ids else set()
    chaves = {(categoria_id, enunciado.strip()) for categoria_id, enunciado in existentes}
    validas = []
    for candidata, dados in preliminares:
        if dados["categoria_id"] not in categorias:
            erros.append(_erro(
                candidata, "categoria_id", "categoria_indisponivel",
                "categoria não encontrada ou indisponível para este modo",
            ))
            continue
        chave = (dados["categoria_id"], dados["enunciado"].strip())
        if chave in chaves:
            erros.append(_erro(
                candidata, "enunciado", "pergunta_duplicada",
                "já existe uma pergunta com este enunciado nesta categoria",
            ))
            continue
        chaves.add(chave)
        validas.append(dados)
    return Auditoria(validas=validas, erros=erros, total=len(candidatas))


def _quota_atual(db: Session, usuario_id: UUID, modo: str) -> int:
    modelo = _modelo(modo)
    return db.scalar(select(func.count(modelo.id)).where(
        modelo.origem == ORIGEM_USUARIO,
        modelo.usuario_id == usuario_id,
        modelo.excluida_em.is_(None),
    )) or 0


def _token_preview(
    usuario_id: UUID, modo: str, formato: str, digest: str, quantidade: int,
    agora: datetime | None = None,
) -> tuple[str, datetime]:
    emitido = agora or datetime.now(timezone.utc)
    expira = emitido + EXPIRACAO_PREVIEW
    token = jwt.encode({
        "sub": str(usuario_id), "iat": emitido, "exp": expira,
        "finalidade": FINALIDADE_TOKEN, "modo": modo, "formato": formato,
        "sha256": digest, "quantidade": quantidade,
        "versao": VERSAO_CONTRATO,
    }, obter_segredo_jwt(), algorithm=JWT_ALGORITHM)
    return token, expira


def _decodificar_token(token: str, usuario_id: UUID) -> dict[str, Any]:
    try:
        dados = jwt.decode(
            token,
            obter_segredo_jwt(),
            algorithms=[JWT_ALGORITHM],
            options={
                "require": [
                    "sub", "iat", "exp", "finalidade", "modo", "formato",
                    "sha256", "quantidade", "versao",
                ]
            },
        )
    except ExpiredSignatureError:
        raise TokenPreviewInvalido("token de preview expirado") from None
    except InvalidTokenError:
        raise TokenPreviewInvalido("token de preview inválido") from None
    esperado = {
        "finalidade": FINALIDADE_TOKEN,
        "versao": VERSAO_CONTRATO,
        "sub": str(usuario_id),
    }
    if any(dados.get(chave) != valor for chave, valor in esperado.items()):
        raise TokenPreviewInvalido("token de preview inválido")
    if dados.get("modo") not in {MODO_QUIZ_CLASSICO, MODO_NEM_A_PATO}:
        raise TokenPreviewInvalido("token de preview inválido")
    if dados.get("formato") not in FORMATOS:
        raise TokenPreviewInvalido("token de preview inválido")
    if (
        not isinstance(dados.get("sha256"), str)
        or re.fullmatch(r"[0-9a-f]{64}", dados["sha256"]) is None
        or isinstance(dados.get("quantidade"), bool)
        or not isinstance(dados.get("quantidade"), int)
        or not 1 <= dados["quantidade"] <= MAX_REGISTROS
    ):
        raise TokenPreviewInvalido("token de preview inválido")
    return dados


class MeuConteudoImportacoesService:
    def validar(
        self, db: Session, usuario_id: UUID, modo: str, nome: str,
        conteudo: bytes,
    ) -> PreviewImportacaoPrivada:
        formato = formato_do_nome(nome)
        candidatas = ler(conteudo, formato, modo)
        auditoria = _auditar(db, usuario_id, modo, formato, candidatas)
        atual = _quota_atual(db, usuario_id, modo)
        novas = len(auditoria.validas)
        quota_ok = atual + novas <= LIMITE_PERGUNTAS_USUARIO_POR_MODO
        erros = list(auditoria.erros)
        pode_confirmar = not erros and quota_ok and novas == auditoria.total
        digest = hashlib.sha256(conteudo).hexdigest()
        token = None
        expira = None
        if pode_confirmar:
            token, expira = _token_preview(
                usuario_id, modo, formato, digest, auditoria.total
            )
        return PreviewImportacaoPrivada(
            modo=modo,
            formato=formato,
            arquivo=ArquivoImportacaoPreview(
                nome=nome, tamanho=len(conteudo), sha256=digest
            ),
            quantidade_recebida=auditoria.total,
            quantidade_valida=len(auditoria.validas),
            quantidade_invalida=len(erros),
            quota=QuotaImportacaoPreview(
                atual=atual, novas=novas, apos_confirmacao=atual + novas,
                limite=LIMITE_PERGUNTAS_USUARIO_POR_MODO,
            ),
            pode_confirmar=pode_confirmar,
            erros=erros,
            preview=[self._serializar(item) for item in auditoria.validas],
            token_preview=token,
            expira_em=expira,
        )

    def confirmar(
        self, db: Session, usuario_id: UUID, token: str, nome: str,
        conteudo: bytes,
    ) -> ResultadoImportacaoPrivada:
        dados_token = _decodificar_token(token, usuario_id)
        formato = formato_do_nome(nome)
        digest = hashlib.sha256(conteudo).hexdigest()
        if (
            formato != dados_token.get("formato")
            or digest != dados_token.get("sha256")
        ):
            raise TokenPreviewInvalido("arquivo diferente do preview validado")
        modo = dados_token["modo"]
        candidatas = ler(conteudo, formato, modo)
        if len(candidatas) != dados_token.get("quantidade"):
            raise TokenPreviewInvalido("arquivo diferente do preview validado")
        try:
            usuario = db.scalar(
                select(Usuario).where(Usuario.id == usuario_id).with_for_update()
            )
            if usuario is None:
                raise TokenPreviewInvalido("token de preview inválido")
            auditoria = _auditar(
                db, usuario_id, modo, formato, candidatas,
                bloquear_categorias=True,
            )
            if auditoria.erros or len(auditoria.validas) != auditoria.total:
                raise ConflitoImportacaoPrivada
            atual = _quota_atual(db, usuario_id, modo)
            if atual + auditoria.total > LIMITE_PERGUNTAS_USUARIO_POR_MODO:
                raise QuotaImportacaoPrivadaAlterada
            if modo == MODO_QUIZ_CLASSICO:
                if db.bind is not None and db.bind.dialect.name == "postgresql":
                    db.execute(text("LOCK TABLE perguntas IN SHARE ROW EXCLUSIVE MODE"))
                proximo = (db.scalar(select(func.max(Pergunta.id))) or 0) + 1
                perguntas = [Pergunta(
                    id=proximo + indice, **dados, origem=ORIGEM_USUARIO,
                    usuario_id=usuario_id, ativa=True,
                ) for indice, dados in enumerate(auditoria.validas)]
            else:
                perguntas = [PerguntaNemPato(
                    **dados, origem=ORIGEM_USUARIO, usuario_id=usuario_id,
                    ativa=True,
                ) for dados in auditoria.validas]
            db.add_all(perguntas)
            db.flush()
            db.commit()
            return ResultadoImportacaoPrivada(
                modo=modo, formato=formato, criadas=len(perguntas)
            )
        except Exception:
            db.rollback()
            raise

    @staticmethod
    def _serializar(dados: dict[str, Any]) -> dict[str, Any]:
        resultado = dict(dados)
        resultado["categoria_id"] = str(resultado["categoria_id"])
        if "alternativa_correta" in resultado:
            resultado["alternativa_correta"] = tuple(ALTERNATIVAS)[
                resultado["alternativa_correta"]
            ]
        return resultado


meu_conteudo_importacoes_service = MeuConteudoImportacoesService()
