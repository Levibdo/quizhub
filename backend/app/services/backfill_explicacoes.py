"""Backfill administrativo; nunca utilizado pelo importador ou pelos endpoints."""

import argparse
from collections import Counter
from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook
from sqlalchemy import select, text, update
from sqlalchemy.engine import Engine

from app.models import Pergunta


DISTRIBUICAO = {"geral": 35, "matematica": 35, "tecnologia": 35}
COLUNAS = ("categoria_id", "enunciado", "explicacao")


def ler_explicacoes(conteudo: bytes) -> dict[tuple[str, str], str]:
    workbook = load_workbook(BytesIO(conteudo), read_only=True, data_only=False)
    try:
        linhas = workbook.active.iter_rows()
        cabecalho = [cell.value for cell in next(linhas, ())]
        if any(cabecalho.count(nome) != 1 for nome in COLUNAS):
            raise ValueError("Colunas obrigatorias ausentes ou repetidas")
        indices = [cabecalho.index(nome) for nome in COLUNAS]
        dados = {}
        quantidade = 0
        for numero, linha in enumerate(linhas, 2):
            quantidade += 1
            valores = []
            for nome, indice in zip(COLUNAS, indices):
                cell = linha[indice]
                valor = cell.value
                if cell.data_type == "f" or not isinstance(valor, str) or not valor.strip():
                    raise ValueError(f"Linha {numero}: {nome} vazio ou texto invalido")
                valores.append(valor)
            categoria, enunciado, explicacao = valores
            if categoria not in DISTRIBUICAO:
                raise ValueError(f"Linha {numero}: categoria invalida")
            # Identidade exata: nao normalizar caixa, espacos ou acentos.
            chave = (categoria, enunciado)
            if chave in dados:
                raise ValueError(f"Linha {numero}: duplicata no XLSX")
            dados[chave] = explicacao.strip()
        if quantidade != 105:
            raise ValueError(f"Esperadas 105 linhas; encontradas {quantidade}")
        if Counter(categoria for categoria, _ in dados) != DISTRIBUICAO:
            raise ValueError("Distribuicao deve ser 35/35/35")
        return dados
    finally:
        workbook.close()


def backfill_explicacoes(engine: Engine, conteudo: bytes, *, aplicar: bool = False) -> int:
    """Valida tudo em transacao propria; aplicar=False realiza somente auditoria.

    A repeticao executa novamente 105 UPDATEs com os mesmos valores.
    Qualquer excecao, inclusive divergencia apos os UPDATEs, desfaz a transacao.
    """
    tabela = Pergunta.__table__
    with engine.begin() as conn:
        dados = ler_explicacoes(conteudo)
        if conn.dialect.name == "postgresql":
            conn.execute(text("LOCK TABLE perguntas IN SHARE ROW EXCLUSIVE MODE"))
        antes = [dict(row) for row in conn.execute(select(tabela)).mappings()]
        chaves = [(row["categoria_id"], row["enunciado"]) for row in antes]
        if len(set(chaves)) != len(chaves):
            raise ValueError("Correspondencia multipla no banco")
        ausentes_banco = set(dados) - set(chaves)
        ausentes_xlsx = set(chaves) - set(dados)
        if ausentes_banco or ausentes_xlsx:
            raise ValueError(
                f"Ausentes no banco: {len(ausentes_banco)}; "
                f"ausentes no XLSX: {len(ausentes_xlsx)}"
            )
        if len(antes) != 105:
            raise ValueError("Exigidas 105 correspondencias")
        if not aplicar:
            return len(antes)
        atualizadas = 0
        for (categoria, enunciado), explicacao in dados.items():
            resultado = conn.execute(
                update(tabela).where(
                    tabela.c.categoria_id == categoria,
                    tabela.c.enunciado == enunciado,
                ).values(explicacao=explicacao)
            )
            if resultado.rowcount != 1:
                raise ValueError("UPDATE deve afetar exatamente uma pergunta")
            atualizadas += resultado.rowcount
        if atualizadas != 105:
            raise ValueError("Exigidos exatamente 105 updates")
        depois = {
            (row["categoria_id"], row["enunciado"]): dict(row)
            for row in conn.execute(select(tabela)).mappings()
        }
        esperado = {
            (row["categoria_id"], row["enunciado"]): {
                **row, "explicacao": dados[(row["categoria_id"], row["enunciado"])]
            }
            for row in antes
        }
        if depois != esperado:
            raise ValueError("Validacao final falhou: campos originais ou explicacoes divergentes")
    return atualizadas


def main() -> None:
    from app.db.session import get_engine

    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("xlsx", type=Path)
    parser.add_argument("--aplicar", action="store_true", help="Efetiva os 105 updates")
    args = parser.parse_args()
    quantidade = backfill_explicacoes(
        get_engine(), args.xlsx.read_bytes(), aplicar=args.aplicar
    )
    print(f"{'Atualizadas' if args.aplicar else 'Correspondencias'}: {quantidade}")


if __name__ == "__main__":
    main()
