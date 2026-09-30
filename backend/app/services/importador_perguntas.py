import csv
import json
from dataclasses import dataclass
from io import BytesIO, StringIO
from pathlib import Path
from zipfile import BadZipFile

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from sqlalchemy import func, select, text
from sqlalchemy.orm import Session

from app.models import Categoria, Pergunta
from app.schemas.perguntas import (
    ErroImportacaoPergunta,
    RelatorioImportacaoPerguntas,
    RelatorioValidacaoPerguntas,
)

COLUNAS_OBRIGATORIAS = (
    "categoria_id", "enunciado", "alternativa_a", "alternativa_b",
    "alternativa_c", "alternativa_d", "alternativa_correta", "explicacao",
)
FORMATOS_SUPORTADOS = {"xlsx", "csv", "json"}


class ArquivoImportacaoInvalido(ValueError):
    pass


@dataclass(frozen=True)
class PerguntaCandidata:
    referencia: int
    dados: dict[str, object]
    erro_estrutura: str | None = None


@dataclass(frozen=True)
class AuditoriaPerguntas:
    formato: str
    total: int
    validas: list[dict[str, object]]
    duplicadas: int
    invalidas: int
    erros: list[ErroImportacaoPergunta]

    def relatorio(self) -> RelatorioValidacaoPerguntas:
        return RelatorioValidacaoPerguntas(
            formato=self.formato, total=self.total, validas=len(self.validas),
            duplicadas=self.duplicadas, invalidas=self.invalidas, erros=self.erros,
        )


def detectar_formato(caminho: str | Path) -> str:
    formato = Path(caminho).suffix.lower().removeprefix(".")
    if formato not in FORMATOS_SUPORTADOS:
        raise ArquivoImportacaoInvalido(
            "extensão não suportada; use .xlsx, .csv ou .json"
        )
    return formato


def _texto_obrigatorio(valor) -> str | None:
    if valor is None:
        return None
    texto_valor = str(valor).strip()
    return texto_valor or None


def _alternativa_correta(valor) -> int | None:
    if isinstance(valor, bool):
        return None
    try:
        numero = int(valor)
    except (TypeError, ValueError):
        return None
    if isinstance(valor, float) and not valor.is_integer():
        return None
    return numero if 0 <= numero <= 3 else None


def _indices_cabecalho(cabecalho_original) -> dict[str, int]:
    if cabecalho_original is None:
        raise ArquivoImportacaoInvalido("arquivo sem cabeçalho")
    cabecalho = [
        str(valor).strip() if valor is not None else ""
        for valor in cabecalho_original
    ]
    ausentes = [coluna for coluna in COLUNAS_OBRIGATORIAS if coluna not in cabecalho]
    if ausentes:
        raise ArquivoImportacaoInvalido(
            "colunas obrigatórias ausentes: " + ", ".join(ausentes)
        )
    return {coluna: cabecalho.index(coluna) for coluna in COLUNAS_OBRIGATORIAS}


def ler_xlsx(conteudo: bytes) -> list[PerguntaCandidata]:
    try:
        workbook = load_workbook(BytesIO(conteudo), read_only=True, data_only=True)
    except (BadZipFile, InvalidFileException, KeyError, OSError) as erro:
        raise ArquivoImportacaoInvalido("arquivo XLSX inválido") from erro
    try:
        linhas = workbook.active.iter_rows(values_only=True)
        indices = _indices_cabecalho(next(linhas, None))
        return [
            PerguntaCandidata(numero, {
                coluna: valores[indice] if indice < len(valores) else None
                for coluna, indice in indices.items()
            })
            for numero, valores in enumerate(linhas, start=2)
        ]
    finally:
        workbook.close()


def ler_csv(conteudo: bytes) -> list[PerguntaCandidata]:
    try:
        texto_csv = conteudo.decode("utf-8-sig")
    except UnicodeDecodeError as erro:
        raise ArquivoImportacaoInvalido("arquivo CSV deve usar UTF-8") from erro
    try:
        leitor = csv.reader(StringIO(texto_csv), delimiter=",", strict=True)
        indices = _indices_cabecalho(next(leitor, None))
        return [
            PerguntaCandidata(numero, {
                coluna: valores[indice] if indice < len(valores) else None
                for coluna, indice in indices.items()
            })
            for numero, valores in enumerate(leitor, start=2)
        ]
    except csv.Error as erro:
        raise ArquivoImportacaoInvalido(f"arquivo CSV inválido: {erro}") from erro


def ler_json(conteudo: bytes) -> list[PerguntaCandidata]:
    try:
        documento = json.loads(conteudo.decode("utf-8-sig"))
    except (UnicodeDecodeError, json.JSONDecodeError) as erro:
        raise ArquivoImportacaoInvalido("JSON malformado ou fora de UTF-8") from erro
    if not isinstance(documento, dict):
        raise ArquivoImportacaoInvalido("JSON deve conter um objeto no nível superior")
    if documento.get("modo") != "quiz_classico":
        raise ArquivoImportacaoInvalido("modo deve ser exatamente 'quiz_classico'")
    categoria_id = documento.get("categoria_id")
    perguntas = documento.get("perguntas")
    if not isinstance(perguntas, list) or not perguntas:
        raise ArquivoImportacaoInvalido("perguntas deve ser uma lista não vazia")
    candidatas = []
    for indice, pergunta in enumerate(perguntas, start=1):
        if not isinstance(pergunta, dict):
            candidatas.append(PerguntaCandidata(
                indice, {}, "pergunta deve ser um objeto"
            ))
            continue
        alternativas = pergunta.get("alternativas")
        erro_estrutura = None
        if not isinstance(alternativas, list) or len(alternativas) != 4:
            erro_estrutura = "alternativas deve conter exatamente quatro itens"
            alternativas = [None, None, None, None]
        candidatas.append(PerguntaCandidata(indice, {
            "categoria_id": categoria_id,
            "enunciado": pergunta.get("enunciado"),
            "alternativa_a": alternativas[0], "alternativa_b": alternativas[1],
            "alternativa_c": alternativas[2], "alternativa_d": alternativas[3],
            "alternativa_correta": pergunta.get("alternativa_correta"),
            "explicacao": pergunta.get("explicacao"),
        }, erro_estrutura))
    return candidatas


LEITORES = {"xlsx": ler_xlsx, "csv": ler_csv, "json": ler_json}


class ImportadorPerguntas:
    @staticmethod
    def ler(conteudo: bytes, formato: str) -> list[PerguntaCandidata]:
        formato_normalizado = formato.lower().removeprefix(".")
        leitor = LEITORES.get(formato_normalizado)
        if leitor is None:
            raise ArquivoImportacaoInvalido(
                "formato não suportado; use xlsx, csv ou json"
            )
        return leitor(conteudo)

    def validar_arquivo(
        self, db: Session, conteudo: bytes, formato: str
    ) -> RelatorioValidacaoPerguntas:
        candidatas = self.ler(conteudo, formato)
        return self._auditar(db, candidatas, formato).relatorio()

    def importar_arquivo(
        self, db: Session, conteudo: bytes, formato: str
    ) -> RelatorioImportacaoPerguntas:
        candidatas = self.ler(conteudo, formato)
        try:
            if db.bind is not None and db.bind.dialect.name == "postgresql":
                db.execute(text("LOCK TABLE perguntas IN SHARE ROW EXCLUSIVE MODE"))
            auditoria = self._auditar(db, candidatas, formato)
            proximo_id = (db.scalar(select(func.max(Pergunta.id))) or 0) + 1
            for dados in auditoria.validas:
                db.add(Pergunta(id=proximo_id, **dados))
                proximo_id += 1
            db.flush()
            db.commit()
            return RelatorioImportacaoPerguntas(
                total=auditoria.total, criadas=len(auditoria.validas),
                falhas=auditoria.duplicadas + auditoria.invalidas,
                erros=auditoria.erros,
            )
        except Exception:
            db.rollback()
            raise

    def importar(self, db: Session, conteudo: bytes) -> RelatorioImportacaoPerguntas:
        """Compatibilidade da API existente, que continua aceitando somente XLSX."""
        return self.importar_arquivo(db, conteudo, "xlsx")

    def _auditar(
        self, db: Session, candidatas: list[PerguntaCandidata], formato: str
    ) -> AuditoriaPerguntas:
        categorias = {
            categoria.id: categoria for categoria in db.scalars(select(Categoria))
        }
        chaves_conhecidas = set(
            db.execute(select(Pergunta.categoria_id, Pergunta.enunciado)).tuples()
        )
        validas = []
        erros = []
        duplicadas = 0
        invalidas = 0
        for candidata in candidatas:
            validada, motivo = self._validar_candidata(candidata, categorias)
            if motivo is not None:
                invalidas += 1
                erros.append(ErroImportacaoPergunta(
                    linha=candidata.referencia, motivo=motivo
                ))
                continue
            chave = (validada["categoria_id"], validada["enunciado"])
            if chave in chaves_conhecidas:
                duplicadas += 1
                erros.append(ErroImportacaoPergunta(
                    linha=candidata.referencia,
                    motivo=("pergunta duplicada: categoria_id e enunciado "
                            "já existentes"),
                ))
                continue
            chaves_conhecidas.add(chave)
            validas.append(validada)
        return AuditoriaPerguntas(
            formato=formato.lower().removeprefix("."), total=len(candidatas),
            validas=validas, duplicadas=duplicadas, invalidas=invalidas,
            erros=erros,
        )

    @staticmethod
    def _validar_candidata(
        candidata: PerguntaCandidata, categorias: dict[str, Categoria]
    ) -> tuple[dict[str, object] | None, str | None]:
        if candidata.erro_estrutura:
            return None, candidata.erro_estrutura
        dados = candidata.dados
        textos = {
            coluna: _texto_obrigatorio(dados.get(coluna))
            for coluna in COLUNAS_OBRIGATORIAS
            if coluna != "alternativa_correta"
        }
        for coluna, valor in textos.items():
            if valor is None:
                return None, f"campo obrigatório vazio: {coluna}"
        valor_alternativa = dados.get("alternativa_correta")
        if valor_alternativa is None or str(valor_alternativa).strip() == "":
            return None, "campo obrigatório vazio: alternativa_correta"
        alternativa = _alternativa_correta(valor_alternativa)
        if alternativa is None:
            return None, "alternativa_correta deve estar entre 0 e 3"
        categoria = categorias.get(textos["categoria_id"])
        if categoria is None:
            return None, "categoria inexistente"
        if not categoria.ativa:
            return None, "categoria inativa"
        return {**textos, "alternativa_correta": alternativa, "ativa": True}, None


importador_perguntas = ImportadorPerguntas()
